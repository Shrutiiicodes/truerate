"""Shared openpyxl formatting for work papers (Arial, navy headers, INR formats)."""
from __future__ import annotations
import datetime as dt
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
NAVY = "1F3864"
INR = '"₹"#,##0.00;[Red]-"₹"#,##0.00;"-"'
INR0 = '"₹"#,##0;[Red]-"₹"#,##0;"-"'
INT = '#,##0;[Red]-#,##0;"-"'
PCT = '0.00%'
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BAND_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FONT = Font(name=FONT, size=10, color="0000FF")


def base_font(bold=False, size=10, color="000000", italic=False):
    return Font(name=FONT, bold=bold, size=size, color=color, italic=italic)


def wp_header(ws, eng: dict, title: str, ref_suffix: str, ncols: int = 8) -> int:
    """Big-4 style header block. Returns the next free row."""
    ws["A1"] = eng["entity"]; ws["A1"].font = base_font(True, 12, NAVY)
    ws["A2"] = f"{eng['subject']} - {title}"; ws["A2"].font = base_font(True, 11)
    ws["A3"] = f"Period: {eng['period_label']}"; ws["A3"].font = base_font(size=9, italic=True)
    today = dt.date.today().isoformat()
    lbl = get_column_letter(max(ncols - 1, 2))
    val = get_column_letter(max(ncols, 3))
    for r, (k, v) in enumerate([("W/P ref", f"{eng['workpaper_ref']}.{ref_suffix}"), ("Prepared by", eng["prepared_by"]),
                                ("Date", today), ("Reviewed by", eng["reviewed_by"])], start=1):
        ws[f"{lbl}{r}"] = k; ws[f"{lbl}{r}"].font = base_font(True, 9)
        ws[f"{val}{r}"] = v; ws[f"{val}{r}"].font = base_font(size=9)
    return 6


def table(ws, df, start_row: int, formats: dict | None = None, widths: dict | None = None, freeze=True):
    """Write a DataFrame as a formatted table. Returns (first_data_row, last_data_row)."""
    formats = formats or {}
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=start_row, column=j, value=col)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
    for i, row in enumerate(df.itertuples(index=False), start=start_row + 1):
        for j, v in enumerate(row, start=1):
            if v is not None and not (isinstance(v, float) and v != v):
                ws.cell(row=i, column=j, value=v)
    last = start_row + len(df)
    for j, col in enumerate(df.columns, start=1):
        fmt = formats.get(col)
        letter = get_column_letter(j)
        for (cell,) in ws.iter_rows(min_row=start_row + 1, max_row=last, min_col=j, max_col=j):
            cell.font = base_font(size=9)
            if fmt:
                cell.number_format = fmt
        ws.column_dimensions[letter].width = (widths or {}).get(col, min(max(len(str(col)) + 2, 12), 45))
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(df.columns))}{last}"
    return start_row + 1, last


def text_block(ws, rows, start_row: int, col_a_width=28, col_b_width=110):
    ws.column_dimensions["A"].width = col_a_width
    ws.column_dimensions["B"].width = col_b_width
    r = start_row
    for k, v in rows:
        a = ws.cell(row=r, column=1, value=k); a.font = base_font(True); a.alignment = Alignment(vertical="top", wrap_text=True)
        b = ws.cell(row=r, column=2, value=v); b.font = base_font(); b.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    return r
